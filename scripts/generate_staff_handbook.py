from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "876Nurses_Staff_User_Handbook_Draft.docx"
LOGO = ROOT / "assets" / "Images" / "Nurses-icon.png"
COVER_IMAGE = ROOT / "assets" / "Images" / "handbook-cover.png"

NAVY = RGBColor(10, 43, 73)
SKY = RGBColor(54, 166, 224)
CYAN = RGBColor(125, 214, 240)
TEAL = RGBColor(19, 124, 148)
SLATE = RGBColor(78, 91, 107)
LIGHT = RGBColor(242, 247, 250)
WHITE = RGBColor(255, 255, 255)

SECTIONS = [
    (
        "01",
        "Welcome",
        "A quick orientation to the 876Nurses platform, your daily workflow, and how the handbook is structured.",
        [
            "What 876Nurses is used for across admin, nurse, and finance workflows.",
            "How to navigate the mobile app confidently from your first login.",
            "Where to find the most common tasks and what to do when something looks wrong.",
        ],
    ),
    (
        "02",
        "Installation and Setup",
        "Device setup, sign-in preparation, required permissions, and first-time configuration steps.",
        [
            "Install the app and confirm device compatibility.",
            "Enable notifications, camera, photos, and location only when required by role.",
            "Confirm you can sign in and reach the correct account dashboard.",
        ],
    ),
    (
        "03",
        "Getting Started",
        "First login, password basics, and the initial checks every staff member should complete.",
        [
            "Review profile details and confirm your role-specific access is correct.",
            "Check unread alerts, pending actions, and the current day view.",
            "Log out and back in once if your access or data looks incomplete.",
        ],
    ),
    (
        "04",
        "Dashboard and Overview",
        "Reading the dashboard, understanding action cards, and locating the main app controls.",
        [
            "Identify summary cards, counters, and quick actions.",
            "Understand where appointments, schedules, and finance items surface first.",
            "Use this section for screenshots of the main dashboard and home navigation.",
        ],
    ),
    (
        "05",
        "Customers",
        "Patient and customer records, booking flow, notes, and follow-up visibility.",
        [
            "Create or update customer records accurately.",
            "Book appointments and verify payment or deposit expectations.",
            "Capture notes, attachments, and status updates clearly.",
        ],
    ),
    (
        "06",
        "Nurses",
        "Shift assignment, schedule visibility, responses, and nurse-specific day-to-day actions.",
        [
            "Review schedules and pending confirmations.",
            "Track nurse status changes, reassignment needs, and completed work.",
            "Use screenshot slots for shift lists, schedule details, and response screens.",
        ],
    ),
    (
        "07",
        "Administrators",
        "Staff management, approvals, operational oversight, and exception handling.",
        [
            "Add staff, review access levels, and confirm sequence-based codes.",
            "Manage approvals, escalations, and operational corrections.",
            "Keep screenshots focused on high-frequency admin actions.",
        ],
    ),
    (
        "08",
        "Accountants",
        "Invoices, payment tracking, deposits, and finance reconciliation workflows.",
        [
            "Review invoice generation and numbering expectations.",
            "Confirm deposit handling, payment status, and payout tracking.",
            "Use screenshot slots for invoice lists, payment details, and reports.",
        ],
    ),
    (
        "09",
        "FAQs",
        "Short answers to the questions staff ask most often during onboarding and daily use.",
        [
            "What to do if the wrong dashboard appears.",
            "How to refresh stale information or missing records.",
            "Where to report issues that need a system-level fix.",
        ],
    ),
    (
        "10",
        "Support Contacts",
        "Escalation routes, internal contacts, and notes for maintaining the handbook over time.",
        [
            "List the person or team to contact for account, schedule, billing, or technical support.",
            "Keep business hours, email, phone, and escalation guidance up to date.",
            "Use this final page as a live admin-owned reference.",
        ],
    ),
]


def set_cell_shading(cell, fill):
    cell_properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell_properties.append(shading)


def set_page_background(section, color_hex):
    section_properties = section._sectPr
    background = section_properties.find(qn("w:background"))
    if background is None:
        background = OxmlElement("w:background")
        section_properties.insert(0, background)
    background.set(qn("w:color"), color_hex)


def style_paragraph(paragraph, size, color, bold=False, name="Aptos"):
    for run in paragraph.runs:
        run.font.name = name
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold


def add_divider(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(" ")
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "7DD6F0")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)
    run.font.size = Pt(2)


def add_cover(document):
    section = document.sections[0]

    if COVER_IMAGE.exists():
        section.top_margin = Inches(0)
        section.bottom_margin = Inches(0)
        section.left_margin = Inches(0)
        section.right_margin = Inches(0)
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(COVER_IMAGE), width=Inches(8.5), height=Inches(11.0))
        new_section = document.add_section(WD_SECTION.NEW_PAGE)
        new_section.top_margin = Inches(0.7)
        new_section.bottom_margin = Inches(0.6)
        new_section.left_margin = Inches(0.65)
        new_section.right_margin = Inches(0.65)
        return

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    hero = document.add_table(rows=1, cols=2)
    hero.alignment = WD_TABLE_ALIGNMENT.CENTER
    hero.allow_autofit = False
    hero.columns[0].width = Inches(4.8)
    hero.columns[1].width = Inches(2.0)

    left = hero.cell(0, 0)
    right = hero.cell(0, 1)
    set_cell_shading(left, "0A2B49")
    set_cell_shading(right, "7DD6F0")
    left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    brand = left.paragraphs[0]
    brand.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = brand.add_run("876NURSES")
    r.font.name = "Aptos Display"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = CYAN

    title = left.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Staff User\nHandbook")
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = WHITE

    subtitle = left.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subrun = subtitle.add_run(
        "Operational guide for administrators, nurses, accountants, and office staff."
    )
    subrun.font.name = "Aptos"
    subrun.font.size = Pt(11)
    subrun.font.color.rgb = WHITE

    note = left.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    noterun = note.add_run(
        "Designed to mirror the Canva handbook structure with clean section numbering and image-ready layouts."
    )
    noterun.font.name = "Aptos"
    noterun.font.size = Pt(9)
    noterun.font.color.rgb = CYAN

    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if LOGO.exists():
        right.paragraphs[0].add_run().add_picture(str(LOGO), width=Inches(1.3))
    badge = right.add_paragraph()
    badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br = badge.add_run("2026 EDITION")
    br.font.name = "Aptos"
    br.font.size = Pt(11)
    br.font.bold = True
    br.font.color.rgb = NAVY

    document.add_paragraph()
    info = document.add_table(rows=1, cols=3)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.allow_autofit = False
    labels = [
        ("Prepared For", "876 Nurses Home Care Services Ltd"),
        ("Document Use", "Internal staff onboarding and daily reference"),
        ("Screenshot Note", "Image slots intentionally left blank for manual insertion"),
    ]
    for cell, (label, value) in zip(info.rows[0].cells, labels):
        set_cell_shading(cell, "F2F7FA")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(label + "\n")
        r1.font.name = "Aptos"
        r1.font.size = Pt(8)
        r1.font.bold = True
        r1.font.color.rgb = TEAL
        r2 = p1.add_run(value)
        r2.font.name = "Aptos"
        r2.font.size = Pt(9)
        r2.font.color.rgb = NAVY

    document.add_paragraph()
    add_divider(document)
    teaser = document.add_table(rows=1, cols=2)
    teaser.alignment = WD_TABLE_ALIGNMENT.CENTER
    teaser.allow_autofit = False
    teaser.columns[0].width = Inches(3.2)
    teaser.columns[1].width = Inches(3.6)

    left_box = teaser.cell(0, 0)
    right_box = teaser.cell(0, 1)
    set_cell_shading(left_box, "137C94")
    set_cell_shading(right_box, "EAF7FC")

    p = left_box.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    rr = p.add_run("Inside this guide\n01-10")
    rr.font.name = "Aptos Display"
    rr.font.size = Pt(20)
    rr.font.bold = True
    rr.font.color.rgb = WHITE

    p2 = right_box.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p2.add_run(
        "Welcome, setup, dashboards, customers, nurses, administrators, accountants, FAQs, and support contact pages are laid out in a visual sequence so screenshots can drop in cleanly later."
    )
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = NAVY

    document.add_page_break()


def add_contents(document):
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run("Table of Contents")
    run.font.name = "Aptos Display"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = intro.add_run(
        "Use this visual index to jump to the workflow area you need. Each section is designed to hold its own screenshots, step notes, and policy reminders."
    )
    r.font.name = "Aptos"
    r.font.size = Pt(10)
    r.font.color.rgb = SLATE

    grid = document.add_table(rows=5, cols=2)
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.allow_autofit = False

    for idx, (number, title, description, _) in enumerate(SECTIONS):
        row = idx // 2
        col = idx % 2
        cell = grid.cell(row, col)
        fill = "F2F7FA" if idx % 2 == 0 else "EAF7FC"
        set_cell_shading(cell, fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        num_p = cell.paragraphs[0]
        num_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        nr = num_p.add_run(number)
        nr.font.name = "Aptos Display"
        nr.font.size = Pt(18)
        nr.font.bold = True
        nr.font.color.rgb = SKY

        title_p = cell.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tr = title_p.add_run(title)
        tr.font.name = "Aptos Display"
        tr.font.size = Pt(13)
        tr.font.bold = True
        tr.font.color.rgb = NAVY

        desc_p = cell.add_paragraph()
        desc_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        dr = desc_p.add_run(description)
        dr.font.name = "Aptos"
        dr.font.size = Pt(9)
        dr.font.color.rgb = SLATE

    document.add_page_break()


def add_section_header(document, number, title, description):
    banner = document.add_table(rows=1, cols=2)
    banner.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner.allow_autofit = False
    banner.columns[0].width = Inches(1.2)
    banner.columns[1].width = Inches(5.8)

    left = banner.cell(0, 0)
    right = banner.cell(0, 1)
    set_cell_shading(left, "0A2B49")
    set_cell_shading(right, "EAF7FC")

    p1 = left.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(number)
    r1.font.name = "Aptos Display"
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = WHITE

    p2 = right.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r2 = p2.add_run(title + "\n")
    r2.font.name = "Aptos Display"
    r2.font.size = Pt(20)
    r2.font.bold = True
    r2.font.color.rgb = NAVY
    r3 = p2.add_run(description)
    r3.font.name = "Aptos"
    r3.font.size = Pt(10)
    r3.font.color.rgb = SLATE


def add_key_points(document, points):
    document.add_paragraph()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = heading.add_run("Key Focus Areas")
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = TEAL

    for point in points:
        paragraph = document.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Inches(0.15)
        paragraph.paragraph_format.space_after = Pt(5)
        bullet = paragraph.add_run("• ")
        bullet.font.name = "Aptos"
        bullet.font.size = Pt(10)
        bullet.font.color.rgb = SKY
        text = paragraph.add_run(point)
        text.font.name = "Aptos"
        text.font.size = Pt(10)
        text.font.color.rgb = NAVY


def add_screenshot_row(document, left_title, right_title):
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.2)

    for cell, label in zip(table.rows[0].cells, [left_title, right_title]):
        set_cell_shading(cell, "FFFFFF")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        title = cell.paragraphs[0]
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = title.add_run(label + "\n")
        r1.font.name = "Aptos"
        r1.font.size = Pt(10)
        r1.font.bold = True
        r1.font.color.rgb = NAVY
        r2 = title.add_run("Screenshot placeholder")
        r2.font.name = "Aptos"
        r2.font.size = Pt(9)
        r2.font.color.rgb = SKY

        for _ in range(8):
            spacer = cell.add_paragraph()
            spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sr = spacer.add_run(" ")
            sr.font.size = Pt(8)

        note = cell.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        nr = note.add_run("Insert cropped app screenshot here")
        nr.font.name = "Aptos"
        nr.font.size = Pt(8)
        nr.font.italic = True
        nr.font.color.rgb = SLATE

        borders = OxmlElement("w:tcBorders")
        for edge in ["top", "left", "bottom", "right"]:
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "10")
            border.set(qn("w:color"), "7DD6F0")
            borders.append(border)
        cell._tc.get_or_add_tcPr().append(borders)


def add_tips_box(document, section_title):
    document.add_paragraph()
    box = document.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.cell(0, 0)
    set_cell_shading(cell, "F2F7FA")
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p.add_run("Notes to Customize\n")
    r1.font.name = "Aptos"
    r1.font.size = Pt(10)
    r1.font.bold = True
    r1.font.color.rgb = TEAL
    r2 = p.add_run(
        f"Use this {section_title.lower()} page to add your final screenshots, role-specific steps, and any office policy notes that staff should see before they start using the app."
    )
    r2.font.name = "Aptos"
    r2.font.size = Pt(9)
    r2.font.color.rgb = NAVY


def add_section_page(document, section_data, last=False):
    number, title, description, points = section_data
    add_section_header(document, number, title, description)
    add_key_points(document, points)
    document.add_paragraph()
    add_screenshot_row(document, f"{title} view A", f"{title} view B")
    add_tips_box(document, title)
    if not last:
        document.add_page_break()


def build_document():
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    add_cover(document)
    add_contents(document)
    for index, section_data in enumerate(SECTIONS):
        add_section_page(document, section_data, last=index == len(SECTIONS) - 1)

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()